import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from pydantic import ValidationError
from sqlalchemy import select, update
from fastapi import HTTPException
from .db import Book, Record, Snapshot, Audit, uid, now
from .schemas import LedgerInput
from .domain import rule_identity

# Header aliases in priority order: the canonical name first, then common Chinese export headers.
# Matching stays exact per alias; the preview shows the chosen column and the user can remap.
ALIASES = {
    'source_id': ['source_id', '来源编号', '记录编号'], 'kind': ['kind', '类型', '记录类型'],
    'business_date': ['business_date', '业务日期', '日期', '凭证日期', '记账日期', '制单日期', '交易日期'],
    'period': ['period', '期间', '会计期间'],
    'currency': ['currency', '币种'], 'counterparty': ['counterparty', '交易对手', '往来单位', '对方单位', '客户', '供应商'],
    'reference': ['reference', '关联编号'], 'order_id': ['order_id', '订单号', '订单编号'],
    'invoice_id': ['invoice_id', '发票号', '发票号码'], 'voucher_id': ['voucher_id', '凭证号', '凭证编号', '凭证字号'],
    'account': ['account', '科目', '科目编码', '科目代码', '会计科目'],
    'debit': ['debit', '借方', '借方金额', '借方发生额'], 'credit': ['credit', '贷方', '贷方金额', '贷方发生额'],
    'amount': ['amount', '金额'], 'description': ['description', '说明', '摘要', '备注'],
    'direction': ['direction', '方向', '借贷方向'],
}
FIELD_LABELS = {'source_id': '来源编号', 'kind': '记录类型', 'business_date': '业务日期', 'period': '会计期间',
                'currency': '币种', 'counterparty': '交易对手', 'reference': '关联编号', 'order_id': '订单号',
                'invoice_id': '发票号', 'voucher_id': '凭证号', 'account': '科目', 'debit': '借方金额',
                'credit': '贷方金额', 'amount': '业务金额', 'description': '摘要', 'direction': '收付方向'}
DATE_FORMAT = re.compile(r'^(\d{4})[/.年-](\d{1,2})[/.月-](\d{1,2})日?$')


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def derive_voucher_fields(mapped: dict, filename: str, row_number: int) -> list[str]:
    """Fill fields that ordinary voucher exports leave out, and return which ones were derived.

    Only structural facts are derived, never business judgments: a line with an account and a debit
    or credit is a journal line; the period is the voucher date's month; the source number is the
    file and row; a blank opposite side of a one-sided journal line is zero. A line with both sides
    blank, an unparseable date or any other missing amount still fails validation.
    """
    derived = []
    date = mapped.get('business_date')
    if isinstance(date, str) and (match := DATE_FORMAT.match(date.strip())):
        normalized = f'{match[1]}-{int(match[2]):02d}-{int(match[3]):02d}'
        if normalized != date:
            mapped['business_date'] = normalized
            derived.append('business_date')
    sides = [not _blank(mapped.get(side)) for side in ('debit', 'credit')]
    if _blank(mapped.get('kind')) and not _blank(mapped.get('account')) and any(sides):
        mapped['kind'] = 'journal'
        derived.append('kind')
    date = mapped.get('business_date')
    if _blank(mapped.get('period')) and isinstance(date, str) and re.match(r'^\d{4}-\d{2}-\d{2}$', date):
        mapped['period'] = mapped['business_date'][:7]
        derived.append('period')
    if _blank(mapped.get('source_id')):
        mapped['source_id'] = f'{Path(filename).name[:160]}#row-{row_number}'
        derived.append('source_id')
    if mapped.get('kind') == 'journal' and sides.count(True) == 1:
        blank = 'credit' if sides[0] else 'debit'
        mapped[blank] = '0'
        derived.append(blank)
    return derived


def _message(err: dict) -> tuple[str, str]:
    field = '.'.join(map(str, err['loc']))
    label = FIELD_LABELS.get(field, field)
    kind, text = err.get('type', ''), err.get('msg', '')
    if kind == 'missing':
        text = f'缺少必填字段“{label}”；请在字段映射中选择对应列'
    elif kind.startswith('date'):
        text = f'“{label}”不是有效日期，应为 YYYY-MM-DD 或 YYYY/M/D'
    elif kind == 'literal_error':
        text = f'“{label}”取值无效：{text.split("Input should be ")[-1]}'
    elif kind == 'string_pattern_mismatch':
        text = f'“{label}”格式不正确' + ('，应为 YYYY-MM' if field == 'period' else '')
    elif text.startswith('Value error, '):
        text = text[len('Value error, '):]
    return field or 'record', text


def digest(data):
    return hashlib.sha256(data).hexdigest()


def unique_fields(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f'JSON 字段重复：{key}')
        result[key] = value
    return result


def validate_headers(headers):
    if not headers or any(not header for header in headers) or len(headers) != len(set(headers)):
        raise ValueError('表头不能为空或重复；请为每一列提供唯一字段名')


def declares_public_example(value):
    """Recognize the reserved public-example envelope, not authenticity by itself."""
    return isinstance(value, dict) and (
        value.get('data_nature') == 'public_standard_example_projection'
        or str(value.get('schema_version', '')).startswith('mof-einv-example-projection-'))


def validate_public_example(value, mapping):
    if not declares_public_example(value):
        return
    from datasets.mof_vouchers import verify_package
    try:
        verify_package(value)
    except ValueError as exc:
        raise ValueError('公开教学样例包与固定原件不一致；请从官方原件重新生成投影，未导入修改内容') from exc
    if mapping and any(key != column for key, column in mapping.items()):
        raise ValueError('公开教学样例使用原有分录字段，不允许重新映射；教学改编须另建明确标注的资料')


def parse_file(data: bytes, filename: str, mapping: dict | None = None):
    if mapping is not None and (not isinstance(mapping, dict) or any(not isinstance(v, str) for v in mapping.values())):
        raise ValueError('字段映射必须是字段名到列名的对象')
    suffix = Path(filename).suffix.lower()
    locations = []
    extracted = []
    if suffix == '.csv':
        text = data.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        validate_headers(reader.fieldnames)
        raw = list(reader)
        if any(None in row for row in raw):
            raise ValueError('CSV 存在多于表头的列，请检查逗号与引号')
    elif suffix == '.json':
        raw = json.loads(data, object_pairs_hook=unique_fields)
        validate_public_example(raw, mapping)
        if isinstance(raw, dict):
            raw = raw.get('records', [])
        if not isinstance(raw, list) or any(not isinstance(r, dict) for r in raw):
            raise ValueError('JSON 必须是记录数组或包含 records 数组的对象')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel 解压大小超过 100 MB')
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(x or '').strip() for x in next(iterator, [])]
        validate_headers(headers)
        raw = []
        for index, values in enumerate(iterator, 2):
            if index > 50002:
                raise ValueError('单次导入最多 50000 行')
            if not any(x is not None for x in values):
                continue
            raw.append({k: (v.isoformat()[:10] if hasattr(v, 'isoformat') else str(v) if v is not None else '')
                        for k, v in zip(headers, values) if k})
            locations.append({'sheet': sheet.title, 'row': index})
        workbook.close()
    elif suffix == '.pdf':
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > 300:
                raise ValueError('单次 PDF 最多 300 页')
            for number, page in enumerate(pdf.pages, 1):
                extracted.append({'page': number, 'text': page.extract_text() or '',
                                  'status': 'parsed' if page.chars else 'ocr_required'})
        return [], [], [], {}, extracted
    else:
        raise ValueError('支持 UTF-8 CSV、JSON、XLSX 和文本 PDF')
    if len(raw) > 50000:
        raise ValueError('单次导入最多 50000 行')
    columns = list(dict.fromkeys(k for row in raw for k in row))
    inferred = {k: next((alias for alias in aliases if alias in columns), k) for k, aliases in ALIASES.items()}
    if mapping:
        unknown = set(mapping) - set(ALIASES)
        if unknown:
            raise ValueError(f'未知目标字段：{sorted(unknown)}')
        inferred.update(mapping)
    rows, errors = [], []
    for index, row in enumerate(raw):
        locator = locations[index] if locations else {'row': index + (2 if suffix == '.csv' else 1)}
        mapped = {key: row[column] for key, column in inferred.items() if column in row and row[column] is not None}
        derived = derive_voucher_fields(mapped, filename, locator['row'])
        if derived:
            locator = {**locator, 'derived': derived}
        try:
            record = LedgerInput.model_validate(mapped).model_dump(mode='json', exclude_none=True)
            rows.append({'record': record, 'locator': locator})
        except ValidationError as exc:
            for err in exc.errors():
                field, message = _message(err)
                errors.append({'row': locator['row'], 'field': field, 'message': message})
    if not raw:
        errors.append({'row': 0, 'field': 'file', 'message': '未找到可导入记录'})
    return rows, errors, columns, inferred, extracted


def policy_for(book):
    return {'framework': book.framework, 'version': book.version, 'period': book.period,
            'execution_rules': rule_identity(),
            'currency': book.currency, 'accounts': book.accounts, 'locked_periods': book.locked_periods}


def make_snapshot(s, book):
    # Serialize snapshot creation with imports and policy edits; re-read after locking.
    s.execute(update(Book).where(Book.id == book.id).values(revision=Book.revision))
    s.refresh(book)
    records = s.scalars(select(Record).where(Record.book_id == book.id).order_by(Record.id)).all()
    records = [r for r in records if r.payload.get('period') == book.period]
    policy = policy_for(book)
    from .config import settings
    from .knowledge import freeze_reference, load_verified_corpus, search
    corpus_file = settings.data_dir / 'knowledge' / 'corpus-v2.json'
    has_corpus = corpus_file.exists() or corpus_file.is_symlink()
    policy['reference_integrity'] = 'not_configured'
    from .reference_selection import apply_choice
    selected_references = apply_choice(s, book, policy)
    if not selected_references and not has_corpus and (corpus_file.parent / 'corpus.json').exists():
        policy['reference_integrity'] = 'legacy_unverified'
    if not selected_references and has_corpus:
        try:
            corpus = load_verified_corpus(corpus_file, source_root=settings.data_dir / 'public-sources')
        except (ValueError, OSError) as exc:
            raise HTTPException(409, '参考语料来源校验失败；未创建包含失败内容的快照') from exc
        passages = search(corpus, '会计差错更正记账凭证审核规则追溯', book.framework, book.period, limit=10)
        policy['reference_passages'] = [freeze_reference(p) for p in passages]
        policy['reference_integrity'] = 'verified'
    value = json.dumps({'records': [(r.id, r.fingerprint) for r in records], 'policy': policy}, sort_keys=True).encode()
    existing = s.scalar(select(Snapshot).where(Snapshot.book_id == book.id,
                                               Snapshot.revision == book.revision, Snapshot.digest == digest(value)))
    if existing:
        return existing
    snap = Snapshot(book_id=book.id, revision=book.revision, policy=policy,
                    record_ids=[r.id for r in records], digest=digest(value))
    s.add(snap)
    s.flush()
    return snap


def snapshot_records(s, snapshot):
    ids = snapshot.record_ids
    found = []
    for i in range(0, len(ids), 800):
        found.extend(s.scalars(select(Record).where(Record.book_id == snapshot.book_id, Record.id.in_(ids[i:i + 800]))).all())
    if len(found) != len(ids):
        raise HTTPException(409, '快照记录缺失，禁止继续调查')
    return [r.payload for r in sorted(found, key=lambda x: x.id)]


def confirm_import(s, batch):
    # Acquire the book write lock before re-reading batch/source IDs; works for SQLite and PostgreSQL.
    s.execute(update(Book).where(Book.id == batch.book_id).values(revision=Book.revision))
    s.refresh(batch)
    if batch.status == 'confirmed':
        return {'snapshot_id': batch.snapshot_id, 'imported': batch.imported, 'duplicate_count': batch.duplicate_count}
    from .ocr_workflow import require_current_review
    require_current_review(s, batch)
    if batch.errors:
        raise HTTPException(422, {'message': '修复或重新映射错误行后才能确认，未静默丢弃任何行', 'errors': batch.errors[:100]})
    book = s.get(Book, batch.book_id)
    if any(row['record']['period'] != book.period for row in batch.rows):
        raise HTTPException(409, '账套期间已改变或资料不属于当前期间，请重新预览导入')
    existing = {(r.kind, r.source_id): r for r in s.scalars(select(Record).where(Record.book_id == book.id))}
    inserted, duplicate = 0, 0
    for row in batch.rows:
        record = row['record']
        fingerprint = digest(json.dumps(record, sort_keys=True, ensure_ascii=False).encode())
        key = record['kind'], record['source_id']
        if key in existing:
            if existing[key].fingerprint != fingerprint:
                raise HTTPException(409, f'来源编号 {key[1]} 内容冲突；原记录未覆盖，请为源系统修订使用新的来源编号')
            duplicate += 1
            continue
        identity, evidence_id = uid(), uid()
        payload = {**record, 'id': identity, 'evidence_id': evidence_id, 'observed_at': now()}
        item = Record(id=identity, book_id=book.id, document_id=batch.document_id, source_id=key[1], kind=key[0],
                      evidence_id=evidence_id, fingerprint=fingerprint, payload=payload, locator=row['locator'])
        s.add(item)
        existing[key] = item
        inserted += 1
    book.revision += 1
    s.flush()
    snapshot = make_snapshot(s, book)
    batch.status, batch.snapshot_id = 'confirmed', snapshot.id
    batch.imported, batch.duplicate_count = inserted, duplicate
    s.add(Audit(book_id=book.id, action='import.confirmed', details={'import_id': batch.id, 'snapshot_id': snapshot.id,
                                                                  'imported': inserted, 'duplicates': duplicate}))
    s.commit()
    return {'snapshot_id': snapshot.id, 'imported': inserted, 'duplicate_count': duplicate}
